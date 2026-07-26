"""Tier-0 text extraction: pull an existing text layer, no OCR.

PyMuPDF (import pymupdf / fitz) ships manylinux wheels (no system deps, runs on
python:3.12-slim), extracts text more accurately than pure-python readers, and
repairs malformed PDFs. It also rasterizes pages — reused in Phase 4 for OCR +
thumbnails. A PDF with no text layer (a scan) yields empty pages -> caller
routes it to OCR. Returns [(page_no, text), ...], page_no 1-based.
"""
import pymupdf


def _clean(s):
    """Strip NUL (0x00) — Postgres text columns reject it and PDFs occasionally
    embed it in a text layer."""
    return s.replace("\x00", "")


def pages_from_pdf(path):
    out = []
    with pymupdf.open(path) as doc:
        for i, page in enumerate(doc):
            out.append((i + 1, _clean(page.get_text().strip())))
    return out


def pages_from_text(path):
    with open(path, encoding="utf-8", errors="replace") as f:
        return [(1, _clean(f.read()))]


# --- Office formats (OOXML + legacy xls) ------------------------------------
# All pure-python (manylinux wheels, no system deps). Heavy imports are lazy so
# the package still imports without them. Legacy binary .doc/.ppt (OLE) have no
# clean pure-python reader and are routed to OCR/no_text instead.

def pages_from_docx(path):
    import docx
    d = docx.Document(path)
    parts = [p.text for p in d.paragraphs]
    for tbl in d.tables:                        # tables are separate from paragraphs
        for row in tbl.rows:
            parts.append("\t".join(c.text for c in row.cells))
    return [(1, _clean("\n".join(t for t in parts if t)))]


def pages_from_pptx(path):
    from pptx import Presentation
    out = []
    for i, slide in enumerate(Presentation(path).slides):
        texts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                texts.append(shape.text_frame.text)
            if shape.has_table:
                for row in shape.table.rows:
                    texts.append("\t".join(c.text for c in row.cells))
        out.append((i + 1, _clean("\n".join(t for t in texts if t))))
    return out


def pages_from_xlsx(path):
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        out = []
        for i, ws in enumerate(wb.worksheets):
            rows = []
            for row in ws.iter_rows(values_only=True):
                cells = [str(c) for c in row if c is not None]
                if cells:
                    rows.append("\t".join(cells))
            body = (f"[{ws.title}]\n" if ws.title else "") + "\n".join(rows)
            out.append((i + 1, _clean(body)))
        return out
    finally:
        wb.close()


def pages_from_xls(path):
    import xlrd
    book = xlrd.open_workbook(path)
    out = []
    for i in range(book.nsheets):
        sh = book.sheet_by_index(i)
        rows = []
        for r in range(sh.nrows):
            cells = [str(sh.cell_value(r, c)) for c in range(sh.ncols)
                     if sh.cell_value(r, c) not in (None, "")]
            if cells:
                rows.append("\t".join(cells))
        out.append((i + 1, _clean("\n".join(rows))))
    return out


# MIME -> (extractor, engine label). Dispatched in jobs._run_text.
OFFICE_EXTRACTORS = {
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
        (pages_from_docx, "docx"),
    "application/vnd.openxmlformats-officedocument.presentationml.presentation":
        (pages_from_pptx, "pptx"),
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet":
        (pages_from_xlsx, "xlsx"),
    "application/vnd.ms-excel": (pages_from_xls, "xls"),
}
